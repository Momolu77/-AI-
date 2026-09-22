import json
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.domain.errors import DomainError
from app.services.files import atomic_write, digest

MIMES = {
    ".txt": {"text/plain"},
    ".md": {"text/plain", "text/markdown"},
    ".html": {"text/html"},
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
}


def parse_file(filename, mime, raw, settings):
    suffix = Path(filename).suffix.lower()
    if suffix not in settings.attachment_types.split(",") or suffix not in MIMES:
        raise ValueError("不支持的文件格式")
    if mime not in MIMES[suffix] | {"application/octet-stream", ""}:
        raise ValueError("文件类型与扩展名不一致")
    if not raw or len(raw) > settings.attachment_max:
        raise ValueError("文件为空或超过大小上限")
    if suffix == ".pdf":
        if not raw.startswith(b"%PDF-"):
            raise ValueError("PDF 文件签名不正确")
        reader = PdfReader(BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 100:
            raise ValueError("不支持加密或超过 100 页的 PDF")
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif suffix == ".docx":
        with zipfile.ZipFile(BytesIO(raw)) as archive:
            if (
                "word/document.xml" not in archive.namelist()
                or sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024
            ):
                raise ValueError("DOCX 结构无效或解压体积超限")
        doc = Document(BytesIO(raw))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [
                "\t".join(cell.text for cell in row.cells)
                for table in doc.tables
                for row in table.rows
            ]
        )
    else:
        text = raw.decode("utf-8-sig")
        if "\x00" in text or any(ord(char) < 32 and char not in "\t\r\n" for char in text):
            raise ValueError("文本文件含二进制控制字符")
    if not text.strip():
        raise ValueError("没有可读取文字；扫描件请先转为可选中文字的文件")
    if len(text) > 200000:
        raise ValueError("解析文本超过 200000 字符")
    return text


class AttachmentService:
    def __init__(self, core):
        self.core = core

    def recover(self):
        c = self.core
        with c.db.transaction() as conn:
            for row in conn.execute("SELECT id,session_id,body FROM attachments").fetchall():
                body = json.loads(row["body"])
                if body["status"] == "PROCESSING":
                    body.update(
                        status="FAILED",
                        error_type="SERVER_RESTARTED",
                        error_message="服务重启，附件处理未完成，请重新上传。",
                    )
                    c.db.put(conn, "attachments", row["id"], row["session_id"], body)
                    c.emit(
                        conn,
                        row["session_id"],
                        "attachment_failed",
                        {
                            "attachment_id": row["id"],
                            "failure_stage": "processing",
                            "error_type": body["error_type"],
                            "error_message": body["error_message"],
                        },
                    )

    def upload(self, sid, attachment_id, filename, mime, raw):
        c = self.core
        filename = Path(filename.replace("\\", "/")).name[:180]
        with c.db.transaction() as conn:
            c.active(conn, sid)
            old = c.db.get(conn, "attachments", attachment_id)
            if old:
                if (
                    old["session_id"] != sid
                    or old["sha256"] != digest(raw)
                    or old["file_name"] != filename
                ):
                    raise DomainError("ID_CONFLICT", "附件编号冲突。", 409)
                return old
            body = {
                "file_name": filename,
                "mime": mime,
                "file_size": len(raw),
                "sha256": digest(raw),
                "status": "PROCESSING",
                "created_at": c.now(),
                "storage_ref": None,
            }
            c.db.put(conn, "attachments", attachment_id, sid, body)
        try:
            text = parse_file(filename, mime, raw, c.settings)
            ref = f"attachments/{sid}/{attachment_id}{Path(filename).suffix.lower()}"
            atomic_write(c.settings.data_dir / ref, raw)
            body.update(status="READY", storage_ref=ref, extracted_text=text)
            event, payload = (
                "file_attached",
                {
                    "attachment_id": attachment_id,
                    "file_name": filename,
                    "file_type": mime,
                    "file_size": len(raw),
                },
            )
        except Exception as exc:
            body.update(
                status="FAILED",
                error_message=str(exc) if isinstance(exc, ValueError) else "文件无法读取或解析",
                error_type=type(exc).__name__,
            )
            event, payload = (
                "attachment_failed",
                {
                    "attachment_id": attachment_id,
                    "failure_stage": "processing",
                    "error_type": body["error_type"],
                    "error_message": body["error_message"],
                },
            )
        with c.db.transaction() as conn:
            c.db.put(conn, "attachments", attachment_id, sid, body)
            c.emit(conn, sid, event, payload)
        return {"id": attachment_id, "session_id": sid, **body}

    def remove(self, sid, attachment_id):
        c = self.core
        with c.db.transaction() as conn:
            c.active(conn, sid)
            body = c.owned(conn, "attachments", attachment_id, sid)
            body["status"] = "REMOVED"
            c.db.put(conn, "attachments", attachment_id, sid, body)
        return body
